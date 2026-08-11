#!/usr/bin/env python3
"""Site Services VC mapper — normalize site intake, geocode, assign Hub/Local/Remote."""

from __future__ import annotations

import argparse
import math
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------------------------------------------------------------------------
# Constants / schemas
# ---------------------------------------------------------------------------

INPUT_SHEET = "VC_Site_Input_Template"
SHEET_ORDER = [
    "Instructions",
    "VC_Site_Input_Template",
    "VC_Mapped_Sites",
    "VC_Summary",
    "Cost_Model_Load",
    "Exceptions",
    "Run_Metadata",
    "Alias_Map",
]

INPUT_HEADERS = [
    "site_id",
    "site_name",
    "address_full",
    "city",
    "state_province",
    "postal_code",
    "country",
    "users",
    "seat_count",
    "tickets_per_day",
    "tickets_per_month",
    "tickets_per_year",
    "region",
    "district",
    "customer_site_type",
    "notes",
]

MAPPED_HEADERS = [
    "site_id",
    "site_name",
    "address_full",
    "city",
    "state_province",
    "postal_code",
    "country",
    "latitude",
    "longitude",
    "geocode_status",
    "users",
    "seat_count",
    "tickets_per_day",
    "tickets_per_month",
    "tickets_per_year",
    "region",
    "district",
    "customer_site_type",
    "mapped_role",
    "is_hub",
    "hub_site_id",
    "hub_site_name",
    "campus_id",
    "distance_to_hub_miles",
    "est_drive_minutes",
    "within_local_range",
    "catchment_tickets_per_day",
    "dispatch_eligible",
    "notes",
    "exception_flags",
]

SUMMARY_HEADERS = [
    "metric",
    "value",
    "notes",
]

COST_MODEL_HEADERS = [
    "Site",
    "Country",
    "Region",
    "City",
    "State",
    "Address",
    "PostalCode",
    "Latitude",
    "Longitude",
    "TicketsYr",
    "Users",
    "Seats",
    "Devices",
    "DepotHub",
    "Customs",
    "CustomerStaffed",
    "CoverageClass",
    "Critical24x7",
    "SpaceNeeded",
    "Zone",
    "Notes",
    "mapped_role",
    "campus_id",
    "tickets_per_day",
    "day1_techs_hint",
]

EXCEPTIONS_HEADERS = [
    "exception_id",
    "site_id",
    "site_name",
    "exception_type",
    "severity",
    "message",
    "suggested_action",
    "timestamp_utc",
]

RUN_METADATA_HEADERS = [
    "key",
    "value",
]

ALIAS_MAP_HEADERS = [
    "canonical_field",
    "alias",
    "notes",
]

HEADER_FILL = PatternFill("solid", fgColor="13294B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MILES_PER_KM = 0.621371192
EARTH_KM = 6371.0


def _here() -> Path:
    return Path(__file__).resolve().parent


def load_yaml(path: Path) -> Dict[str, Any]:
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def haversine_miles(lat1, lon1, lat2, lon2) -> float:
    if any(pd.isna(x) for x in (lat1, lon1, lat2, lon2)):
        return 9e9
    a, b, c, e = map(math.radians, [float(lat1), float(lon1), float(lat2), float(lon2)])
    km = 2 * EARTH_KM * math.asin(
        math.sqrt(math.sin((c - a) / 2) ** 2 + math.cos(a) * math.cos(c) * math.sin((e - b) / 2) ** 2)
    )
    return km * MILES_PER_KM


def est_drive_minutes(miles: float, speed_mph: float) -> float:
    if miles >= 9e9 or speed_mph <= 0:
        return 9e9
    return (miles / speed_mph) * 60.0


def within_local_range(miles: float, rules: Dict[str, Any]) -> bool:
    """Inclusive OR: miles ≤ cluster_radius OR est drive ≤ max_drive_minutes."""
    if miles >= 9e9:
        return False
    radius = float(rules.get("cluster_radius_miles", 25.0))
    max_min = float(rules.get("max_drive_minutes", 60.0))
    speed = float(rules.get("avg_drive_speed_mph", 40.0))
    # enable_drive_time_logic gates live routing only; haversine proxy always available
    minutes = est_drive_minutes(miles, speed)
    return miles <= radius or minutes <= max_min


def _norm_header(s: Any) -> str:
    return str(s or "").strip().lower().replace("_", " ").replace("-", " ")


def build_alias_lookup(aliases: Dict[str, Any]) -> Dict[str, str]:
    """Map normalized alias → canonical field name."""
    lookup: Dict[str, str] = {}
    for canonical, alist in (aliases or {}).items():
        lookup[_norm_header(canonical)] = canonical
        if isinstance(alist, list):
            for a in alist:
                lookup[_norm_header(a)] = canonical
    return lookup


def rename_columns(df: pd.DataFrame, aliases: Dict[str, Any]) -> pd.DataFrame:
    lookup = build_alias_lookup(aliases)
    rename: Dict[str, str] = {}
    for col in df.columns:
        key = _norm_header(col)
        if key in lookup:
            rename[col] = lookup[key]
    out = df.rename(columns=rename)
    # Drop duplicate canonicals keeping first
    out = out.loc[:, ~out.columns.duplicated()]
    return out


def to_float(val: Any, default: float = 0.0) -> float:
    if val is None or (isinstance(val, float) and math.isnan(val)):
        return default
    if isinstance(val, str) and not val.strip():
        return default
    try:
        return float(val)
    except (TypeError, ValueError):
        return default


def normalize_tickets(row: pd.Series, working_days: float, months: float) -> Tuple[float, float, float]:
    tpd = to_float(row.get("tickets_per_day"), float("nan"))
    tpm = to_float(row.get("tickets_per_month"), float("nan"))
    tpy = to_float(row.get("tickets_per_year"), float("nan"))

    if not math.isnan(tpd) and tpd > 0:
        day = tpd
    elif not math.isnan(tpm) and tpm > 0:
        day = tpm * months / working_days
    elif not math.isnan(tpy) and tpy > 0:
        day = tpy / working_days
    else:
        day = 0.0

    month = day * working_days / months if months else 0.0
    year = day * working_days
    return round(day, 6), round(month, 4), round(year, 2)


def priority_score(row: pd.Series, priority: List[str]) -> Tuple:
    scores = []
    for field in priority:
        scores.append(to_float(row.get(field), 0.0))
    return tuple(scores)


def read_input_sites(path: Path, aliases: Dict[str, Any]) -> pd.DataFrame:
    xl = pd.ExcelFile(path)
    sheet = INPUT_SHEET if INPUT_SHEET in xl.sheet_names else xl.sheet_names[0]
    df = pd.read_excel(path, sheet_name=sheet)
    df = rename_columns(df, aliases)
    # Ensure all input columns exist
    for h in INPUT_HEADERS + ["latitude", "longitude"]:
        if h not in df.columns:
            df[h] = np.nan
    # Drop fully empty rows
    keep = df[INPUT_HEADERS].apply(lambda r: any(str(v).strip() for v in r if pd.notna(v)), axis=1)
    df = df.loc[keep].copy()
    df.reset_index(drop=True, inplace=True)
    return df


def geocode_sites(
    df: pd.DataFrame,
    exceptions: List[Dict[str, Any]],
    user_agent: str = "site-services-vc-mapper/1.0",
    min_delay_sec: float = 1.1,
) -> pd.DataFrame:
    """Fill missing lat/lon via Nominatim. Respect ~1 req/sec rate limit."""
    from geopy.exc import GeocoderServiceError, GeocoderTimedOut
    from geopy.geocoders import Nominatim

    geolocator = Nominatim(user_agent=user_agent, timeout=20)
    lats: List[Any] = []
    lons: List[Any] = []
    statuses: List[str] = []
    last_call = 0.0

    for idx, row in df.iterrows():
        lat = row.get("latitude")
        lon = row.get("longitude")
        if pd.notna(lat) and pd.notna(lon) and str(lat).strip() and str(lon).strip():
            lats.append(float(lat))
            lons.append(float(lon))
            statuses.append("provided")
            continue

        def _clean(val: Any) -> str:
            s = str(val or "").strip()
            if not s or s.lower() in {"nan", "none", "n/a", "na", "twin", "#n/a"}:
                return ""
            # Strip placeholder prefixes like "N/A - TWIN"
            low = s.lower()
            if low.startswith("n/a"):
                return ""
            return s

        address = _clean(row.get("address_full"))
        city = _clean(row.get("city"))
        state = _clean(row.get("state_province"))
        postal = _clean(row.get("postal_code"))
        country = _clean(row.get("country"))
        # Progressive fallbacks: full → city/state/postal/country → city/country
        # Do NOT fall back to country-only (centroids poison catchment).
        queries = []
        for parts in (
            [address, city, state, postal, country],
            [city, state, postal, country],
            [city, country],
        ):
            q = ", ".join(p for p in parts if p)
            if q and q not in queries:
                queries.append(q)

        if not queries:
            lats.append(np.nan)
            lons.append(np.nan)
            statuses.append("missing_address")
            exceptions.append(
                _exc(
                    len(exceptions) + 1,
                    row.get("site_id"),
                    row.get("site_name"),
                    "geocode_missing_address",
                    "high",
                    "No address components available to geocode.",
                    "Add address_full / city / country and rerun.",
                )
            )
            continue

        loc = None
        used_query = ""
        last_error = None
        for query in queries:
            elapsed = time.time() - last_call
            if elapsed < min_delay_sec:
                time.sleep(min_delay_sec - elapsed)
            try:
                loc = geolocator.geocode(query)
                last_call = time.time()
                used_query = query
                if loc is not None:
                    break
            except (GeocoderTimedOut, GeocoderServiceError, Exception) as exc:
                last_call = time.time()
                last_error = exc
                loc = None

        if loc is None:
            lats.append(np.nan)
            lons.append(np.nan)
            statuses.append("geocode_failed")
            msg = f"Nominatim returned no result for: {queries[0]}"
            if last_error:
                msg = f"Geocode error for '{queries[0]}': {last_error}"
            exceptions.append(
                _exc(
                    len(exceptions) + 1,
                    row.get("site_id"),
                    row.get("site_name"),
                    "geocode_failed",
                    "high",
                    msg,
                    "Correct the address or supply latitude/longitude manually.",
                )
            )
        else:
            lats.append(float(loc.latitude))
            lons.append(float(loc.longitude))
            # Mark coarse city/country hits so reviewers can spot weak placements
            if used_query == queries[0]:
                statuses.append("geocoded")
            else:
                statuses.append("geocoded_fallback")
                exceptions.append(
                    _exc(
                        len(exceptions) + 1,
                        row.get("site_id"),
                        row.get("site_name"),
                        "geocode_coarse",
                        "low",
                        f"Used fallback geocode query: {used_query}",
                        "Confirm city-level placement is acceptable for ROM.",
                    )
                )

    out = df.copy()
    out["latitude"] = lats
    out["longitude"] = lons
    out["geocode_status"] = statuses
    return out


def _exc(
    eid: int,
    site_id: Any,
    site_name: Any,
    etype: str,
    severity: str,
    message: str,
    action: str,
) -> Dict[str, Any]:
    return {
        "exception_id": eid,
        "site_id": "" if site_id is None or (isinstance(site_id, float) and math.isnan(site_id)) else site_id,
        "site_name": "" if site_name is None or (isinstance(site_name, float) and math.isnan(site_name)) else site_name,
        "exception_type": etype,
        "severity": severity,
        "message": message,
        "suggested_action": action,
        "timestamp_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


def assign_roles(df: pd.DataFrame, rules: Dict[str, Any], exceptions: List[Dict[str, Any]]) -> pd.DataFrame:
    """
    Catchment-aligned hub / local / remote assignment, processed per country
    (aligned with engine virtual-campus grouping).

    Within each country:
    1. Primary hub maximizes catchment within local range; priority tickets → users → seats.
    2. Within local range of any Staffed campus → always Local (never Remote).
    3. Remote only when tpd < remote_threshold AND outside local range.
    4. Outside local range AND tpd ≥ remote_threshold → Staffed hub/campus.
    """
    hub_threshold = float(rules.get("hub_threshold_tickets_per_day", 3.0))
    remote_threshold = float(rules.get("remote_threshold_tickets_per_day", 1.2))
    priority = list(rules.get("hub_selection_priority") or ["tickets_per_day", "users", "seat_count"])
    enable_dispatch = bool(rules.get("enable_dispatch_override", True))
    speed = float(rules.get("avg_drive_speed_mph", 40.0))

    work = df.copy()
    n = len(work)
    roles = ["Remote"] * n
    is_hub = [False] * n
    hub_ids = [""] * n
    hub_names = [""] * n
    campus_ids = [""] * n
    dist_mi: List[float] = [9e9] * n
    drive_min: List[float] = [9e9] * n
    within = [False] * n
    catchment = [0.0] * n
    dispatch = [False] * n
    flags = [""] * n

    campus_counter = 0

    def sort_key(i: int) -> Tuple:
        row = work.iloc[i]
        tpd = to_float(row["tickets_per_day"], 0.0)
        qualifies = 1 if tpd >= hub_threshold else 0
        return (qualifies, *priority_score(row, priority))

    country_keys = []
    for i in range(n):
        c = str(work.iloc[i].get("country") or "").strip()
        country_keys.append(c if c and c.lower() != "nan" else "_GLOBAL_")
    work["_country_key"] = country_keys

    for country, idxs in work.groupby("_country_key", sort=False).groups.items():
        idxs = list(idxs)
        geo_idx = [
            i for i in idxs
            if pd.notna(work.iloc[i]["latitude"]) and pd.notna(work.iloc[i]["longitude"])
        ]
        no_geo = [i for i in idxs if i not in geo_idx]

        for i in no_geo:
            tpd = to_float(work.iloc[i]["tickets_per_day"], 0.0)
            if tpd >= remote_threshold:
                roles[i] = "Staffed"
                is_hub[i] = tpd >= hub_threshold
                campus_counter += 1
                sid = work.iloc[i].get("site_id") or i + 1
                campus_ids[i] = f"CAMPUS-{campus_counter}-{sid}"
                hub_ids[i] = str(work.iloc[i].get("site_id") or "")
                hub_names[i] = str(work.iloc[i].get("site_name") or "")
                dist_mi[i] = 0.0
                drive_min[i] = 0.0
                flags[i] = "no_coords_promoted_staffed"
                exceptions.append(
                    _exc(
                        len(exceptions) + 1,
                        work.iloc[i].get("site_id"),
                        work.iloc[i].get("site_name"),
                        "no_coordinates",
                        "medium",
                        "No coordinates; promoted to Staffed because tpd ≥ remote threshold.",
                        "Geocode or supply lat/lon to refine catchment.",
                    )
                )
            else:
                roles[i] = "Remote"
                dispatch[i] = enable_dispatch
                campus_ids[i] = "REMOTE"
                flags[i] = "no_coords_remote"
                exceptions.append(
                    _exc(
                        len(exceptions) + 1,
                        work.iloc[i].get("site_id"),
                        work.iloc[i].get("site_name"),
                        "no_coordinates",
                        "medium",
                        "No coordinates; classified Remote (tpd below remote threshold).",
                        "Geocode or supply lat/lon; confirm dispatch coverage.",
                    )
                )

        if not geo_idx:
            continue

        ranked = sorted(geo_idx, key=sort_key, reverse=True)
        best_hub = None
        best_catch = -1.0
        best_pri: Tuple = tuple()
        for i in ranked:
            row = work.iloc[i]
            catch = 0.0
            for j in geo_idx:
                mi = haversine_miles(
                    row["latitude"], row["longitude"],
                    work.iloc[j]["latitude"], work.iloc[j]["longitude"],
                )
                if within_local_range(mi, rules):
                    catch += to_float(work.iloc[j]["tickets_per_day"], 0.0)
            pri = priority_score(row, priority)
            if catch > best_catch or (catch == best_catch and pri > best_pri):
                best_hub, best_catch, best_pri = i, catch, pri

        staffed: set = set()
        if best_hub is not None:
            staffed.add(best_hub)
            catchment[best_hub] = best_catch

        candidates = sorted(
            [i for i in geo_idx if i not in staffed],
            key=lambda i: priority_score(work.iloc[i], priority),
            reverse=True,
        )
        changed = True
        while changed:
            changed = False
            for i in candidates:
                if i in staffed:
                    continue
                lat, lon = work.iloc[i]["latitude"], work.iloc[i]["longitude"]
                near = any(
                    within_local_range(
                        haversine_miles(lat, lon, work.iloc[s]["latitude"], work.iloc[s]["longitude"]),
                        rules,
                    )
                    for s in staffed
                )
                if near:
                    continue
                if to_float(work.iloc[i]["tickets_per_day"], 0.0) >= remote_threshold:
                    staffed.add(i)
                    changed = True

        campus_for_hub: Dict[int, str] = {}
        for s in sorted(staffed, key=lambda i: priority_score(work.iloc[i], priority), reverse=True):
            campus_counter += 1
            sid = work.iloc[s].get("site_id") or f"H{s+1}"
            campus_for_hub[s] = f"CAMPUS-{campus_counter}-{sid}"

        staffed_list = list(staffed)
        for i in geo_idx:
            if i in staffed:
                roles[i] = "Staffed"
                is_hub[i] = True
                hub_ids[i] = str(work.iloc[i].get("site_id") or "")
                hub_names[i] = str(work.iloc[i].get("site_name") or "")
                campus_ids[i] = campus_for_hub[i]
                dist_mi[i] = 0.0
                drive_min[i] = 0.0
                within[i] = True
                catch = 0.0
                for j in geo_idx:
                    mi = haversine_miles(
                        work.iloc[i]["latitude"], work.iloc[i]["longitude"],
                        work.iloc[j]["latitude"], work.iloc[j]["longitude"],
                    )
                    if within_local_range(mi, rules):
                        catch += to_float(work.iloc[j]["tickets_per_day"], 0.0)
                catchment[i] = round(catch, 4)
                if to_float(work.iloc[i]["tickets_per_day"], 0.0) < hub_threshold:
                    flags[i] = "staffed_below_hub_threshold"
                continue

            lat, lon = work.iloc[i]["latitude"], work.iloc[i]["longitude"]
            best_d = 9e9
            best_s = None
            for s in staffed_list:
                mi = haversine_miles(lat, lon, work.iloc[s]["latitude"], work.iloc[s]["longitude"])
                if mi < best_d:
                    best_d = mi
                    best_s = s
            dist_mi[i] = round(best_d, 2) if best_d < 9e9 else 9e9
            drive_min[i] = round(est_drive_minutes(best_d, speed), 1) if best_d < 9e9 else 9e9
            local = within_local_range(best_d, rules) if best_s is not None else False
            within[i] = local
            tpd = to_float(work.iloc[i]["tickets_per_day"], 0.0)

            if local:
                roles[i] = "Local"
                hub_ids[i] = str(work.iloc[best_s].get("site_id") or "")
                hub_names[i] = str(work.iloc[best_s].get("site_name") or "")
                campus_ids[i] = campus_for_hub[best_s]
                catchment[i] = catchment[best_s]
            elif tpd >= remote_threshold:
                roles[i] = "Staffed"
                is_hub[i] = True
                campus_counter += 1
                sid = work.iloc[i].get("site_id") or f"H{i+1}"
                campus_ids[i] = f"CAMPUS-{campus_counter}-{sid}"
                hub_ids[i] = str(sid)
                hub_names[i] = str(work.iloc[i].get("site_name") or "")
                dist_mi[i] = 0.0
                drive_min[i] = 0.0
                staffed.add(i)
                campus_for_hub[i] = campus_ids[i]
                staffed_list = list(staffed)
                flags[i] = "promoted_far_high_volume"
            else:
                roles[i] = "Remote"
                dispatch[i] = enable_dispatch
                campus_ids[i] = "REMOTE"
                if best_s is not None:
                    hub_ids[i] = str(work.iloc[best_s].get("site_id") or "")
                    hub_names[i] = str(work.iloc[best_s].get("site_name") or "")
                flags[i] = "remote_dispatch"

        if not staffed and geo_idx:
            for i in geo_idx:
                roles[i] = "Remote"
                dispatch[i] = enable_dispatch
                campus_ids[i] = "REMOTE"
                flags[i] = "all_remote_low_volume"

    work.drop(columns=["_country_key"], inplace=True, errors="ignore")
    work["mapped_role"] = roles
    work["is_hub"] = is_hub
    work["hub_site_id"] = hub_ids
    work["hub_site_name"] = hub_names
    work["campus_id"] = campus_ids
    work["distance_to_hub_miles"] = [(None if d >= 9e9 else d) for d in dist_mi]
    work["est_drive_minutes"] = [(None if d >= 9e9 else d) for d in drive_min]
    work["within_local_range"] = within
    work["catchment_tickets_per_day"] = catchment
    work["dispatch_eligible"] = dispatch
    work["exception_flags"] = flags
    return work


def coverage_class(role: str) -> str:
    if role == "Staffed":
        return "Dedicated business hours"
    if role == "Local":
        return "Local drive-range"
    return "Dispatch / remote"


def day1_techs_hint(tpd: float, role: str) -> float:
    """Rough Day-1 tech hint for cost-model load (not a full staffing model)."""
    if role == "Remote":
        return 0.0
    # ~8 tickets/day productive capacity baseline; floor 1 for staffed hubs
    if tpd <= 0:
        return 1.0 if role == "Staffed" else 0.0
    raw = tpd / 8.0
    if role == "Staffed":
        return max(1.0, round(raw, 2))
    return round(raw, 2)


def build_summary(mapped: pd.DataFrame, rules: Dict[str, Any]) -> pd.DataFrame:
    role_counts = mapped["mapped_role"].value_counts().to_dict()
    hubs = int(mapped["is_hub"].fillna(False).astype(bool).sum())
    rows = [
        ("site_count", len(mapped), "Total sites processed"),
        ("staffed_count", int(role_counts.get("Staffed", 0)), "Staffed hub / campus sites"),
        ("local_count", int(role_counts.get("Local", 0)), "Local (within catchment of a campus)"),
        ("remote_count", int(role_counts.get("Remote", 0)), "Remote / dispatch"),
        ("hub_count", hubs, "Sites flagged is_hub"),
        ("total_tickets_per_day", round(mapped["tickets_per_day"].sum(), 4), "Sum of normalized tpd"),
        ("total_tickets_per_year", round(mapped["tickets_per_year"].sum(), 2), "Sum of normalized tpy"),
        ("total_users", int(mapped["users"].fillna(0).sum()), "Sum of users"),
        ("hub_threshold_tpd", rules.get("hub_threshold_tickets_per_day"), "From default_rules.yaml"),
        ("remote_threshold_tpd", rules.get("remote_threshold_tickets_per_day"), "From default_rules.yaml"),
        ("cluster_radius_miles", rules.get("cluster_radius_miles"), "Local range miles limb"),
        ("max_drive_minutes", rules.get("max_drive_minutes"), "Local range drive limb (proxy)"),
        ("rules_version", rules.get("rules_version", "1.0.0"), "Rule pack version"),
    ]
    return pd.DataFrame(rows, columns=SUMMARY_HEADERS)


def build_cost_model_load(mapped: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in mapped.iterrows():
        role = str(r.get("mapped_role") or "Remote")
        hub_name = r.get("hub_site_name") or ""
        depot = hub_name if role == "Local" else (r.get("site_name") if role == "Staffed" else "")
        rows.append(
            {
                "Site": r.get("site_name") or r.get("site_id"),
                "Country": r.get("country"),
                "Region": r.get("region"),
                "City": r.get("city"),
                "State": r.get("state_province"),
                "Address": r.get("address_full"),
                "PostalCode": r.get("postal_code"),
                "Latitude": r.get("latitude"),
                "Longitude": r.get("longitude"),
                "TicketsYr": r.get("tickets_per_year"),
                "Users": r.get("users"),
                "Seats": r.get("seat_count"),
                "Devices": np.nan,
                "DepotHub": depot,
                "Customs": "",
                "CustomerStaffed": "Y" if role == "Staffed" else "",
                "CoverageClass": coverage_class(role),
                "Critical24x7": "",
                "SpaceNeeded": "Y" if role == "Staffed" else "",
                "Zone": r.get("district") or r.get("region"),
                "Notes": r.get("notes"),
                "mapped_role": role,
                "campus_id": r.get("campus_id"),
                "tickets_per_day": r.get("tickets_per_day"),
                "day1_techs_hint": day1_techs_hint(to_float(r.get("tickets_per_day"), 0.0), role),
            }
        )
    return pd.DataFrame(rows, columns=COST_MODEL_HEADERS)


def build_alias_map_rows(aliases: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    for canonical, alist in (aliases or {}).items():
        rows.append({"canonical_field": canonical, "alias": canonical, "notes": "canonical"})
        if isinstance(alist, list):
            for a in alist:
                rows.append({"canonical_field": canonical, "alias": a, "notes": "alias"})
    return pd.DataFrame(rows, columns=ALIAS_MAP_HEADERS)


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, max_width: int = 36) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for cell in col[:80]:
            if cell.value is not None:
                width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def write_dataframe(ws, df: pd.DataFrame, headers: List[str]) -> None:
    for j, h in enumerate(headers, 1):
        ws.cell(1, j, h)
    style_header(ws)
    for i, row in enumerate(df.itertuples(index=False), 2):
        mapping = row._asdict() if hasattr(row, "_asdict") else None
        for j, h in enumerate(headers, 1):
            if mapping is not None and h in mapping:
                val = mapping[h]
            elif h in df.columns:
                val = df.iloc[i - 2][h]
            else:
                val = None
            if isinstance(val, (np.floating, float)) and (pd.isna(val) or (isinstance(val, float) and math.isnan(val))):
                val = None
            if isinstance(val, (np.integer,)):
                val = int(val)
            if isinstance(val, (np.bool_,)):
                val = bool(val)
            ws.cell(i, j, val)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    autosize(ws)


def write_instructions(ws, rules: Dict[str, Any]) -> None:
    lines = [
        ("Purpose", "Canonical Site Services VC intake and mapping workbook."),
        ("How to use", "1) Fill VC_Site_Input_Template. 2) Run vc_mapper.py. 3) Review VC_Mapped_Sites, VC_Summary, Exceptions. 4) Feed Cost_Model_Load into the site suite cost model."),
        ("Minimum required fields", "site_id (or site_name), country, and at least one of: tickets_per_day / tickets_per_month / tickets_per_year (or users/seats for later sizing). Address fields strongly recommended for geocoding."),
        ("Ticket rules", f"Normalize using working_days_per_year={rules.get('working_days_per_year', 252)} and months_per_year={rules.get('months_per_year', 12)}. Priority: day → month → year."),
        ("Population rules", "Near a Staffed campus (≤ cluster radius OR ≤ max drive minutes proxy) → always Local. Remote only if far AND tickets/day < remote threshold. Far AND ≥ remote threshold → Staffed campus."),
        ("Hub selection", f"Priority: {', '.join(rules.get('hub_selection_priority') or [])}. Hub threshold tpd: {rules.get('hub_threshold_tickets_per_day')}. Primary hub maximizes catchment within local range."),
        ("Current thresholds", f"hub_threshold={rules.get('hub_threshold_tickets_per_day')} tpd; remote_threshold={rules.get('remote_threshold_tickets_per_day')} tpd; cluster_radius_miles={rules.get('cluster_radius_miles')}; max_drive_minutes={rules.get('max_drive_minutes')}; avg_drive_speed_mph={rules.get('avg_drive_speed_mph', 40)}."),
        ("Output tabs", "VC_Mapped_Sites (roles), VC_Summary (counts), Cost_Model_Load (engine-aligned), Exceptions, Run_Metadata, Alias_Map."),
        ("Geocoding", "If latitude/longitude missing, Nominatim geocoding is attempted (~1 req/sec). Failures are logged on Exceptions."),
        ("Overrides", "Do not silently invent demand. Document reviewer overrides outside this file; rerun mapper after source fixes when possible."),
        ("Rules version", str(rules.get("rules_version", "1.0.0"))),
    ]
    ws.cell(1, 1, "section")
    ws.cell(1, 2, "content")
    style_header(ws)
    for i, (section, content) in enumerate(lines, 2):
        ws.cell(i, 1, section)
        ws.cell(i, 2, content)
        ws.cell(i, 2).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:B{len(lines)+1}"


def preserve_input_sheet(wb_out: Workbook, input_path: Path, aliases: Dict[str, Any]) -> None:
    """Copy input rows into VC_Site_Input_Template on the output workbook."""
    df = read_input_sites(input_path, aliases)
    ws = wb_out.create_sheet(INPUT_SHEET)
    write_dataframe(ws, df[INPUT_HEADERS], INPUT_HEADERS)
    # Name table if rows exist
    end_row = max(2, len(df) + 1)
    end_col = get_column_letter(len(INPUT_HEADERS))
    try:
        table = Table(displayName="tbl_vc_site_input", ref=f"A1:{end_col}{end_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    except Exception:
        pass


def write_output_workbook(
    output_path: Path,
    input_path: Path,
    mapped: pd.DataFrame,
    summary: pd.DataFrame,
    cost_load: pd.DataFrame,
    exceptions: List[Dict[str, Any]],
    metadata: List[Tuple[str, Any]],
    aliases: Dict[str, Any],
    rules: Dict[str, Any],
) -> None:
    wb = Workbook()
    # Remove default; rebuild in exact order
    default = wb.active
    wb.remove(default)

    ws_inst = wb.create_sheet("Instructions")
    write_instructions(ws_inst, rules)

    preserve_input_sheet(wb, input_path, aliases)

    ws_map = wb.create_sheet("VC_Mapped_Sites")
    mapped_out = mapped.reindex(columns=[c for c in MAPPED_HEADERS if c in mapped.columns or c in MAPPED_HEADERS])
    for h in MAPPED_HEADERS:
        if h not in mapped_out.columns:
            mapped_out[h] = np.nan
    mapped_out = mapped_out[MAPPED_HEADERS]
    write_dataframe(ws_map, mapped_out, MAPPED_HEADERS)

    ws_sum = wb.create_sheet("VC_Summary")
    write_dataframe(ws_sum, summary, SUMMARY_HEADERS)

    ws_cost = wb.create_sheet("Cost_Model_Load")
    write_dataframe(ws_cost, cost_load, COST_MODEL_HEADERS)

    exc_df = pd.DataFrame(exceptions, columns=EXCEPTIONS_HEADERS) if exceptions else pd.DataFrame(columns=EXCEPTIONS_HEADERS)
    ws_exc = wb.create_sheet("Exceptions")
    write_dataframe(ws_exc, exc_df, EXCEPTIONS_HEADERS)

    meta_df = pd.DataFrame(metadata, columns=RUN_METADATA_HEADERS)
    ws_meta = wb.create_sheet("Run_Metadata")
    write_dataframe(ws_meta, meta_df, RUN_METADATA_HEADERS)

    alias_df = build_alias_map_rows(aliases)
    ws_alias = wb.create_sheet("Alias_Map")
    write_dataframe(ws_alias, alias_df, ALIAS_MAP_HEADERS)

    # Ensure exact order
    for name in SHEET_ORDER:
        if name in wb.sheetnames:
            wb.move_sheet(name, offset=-len(wb.sheetnames))
    # move_sheet offsets are awkward; rebuild order explicitly
    ordered = [wb[name] for name in SHEET_ORDER if name in wb.sheetnames]
    wb._sheets = ordered  # noqa: SLF001 — intentional tab order

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def run_mapper(
    input_path: Path,
    output_path: Path,
    operator_name: str,
    rules_path: Optional[Path] = None,
    aliases_path: Optional[Path] = None,
    skip_geocode: bool = False,
) -> Dict[str, Any]:
    started = datetime.now(timezone.utc)
    base = _here()
    rules_path = rules_path or (base / "default_rules.yaml")
    aliases_path = aliases_path or (base / "column_aliases.yaml")
    rules = load_yaml(rules_path)
    aliases = load_yaml(aliases_path)

    working_days = float(rules.get("working_days_per_year", 252))
    months = float(rules.get("months_per_year", 12))

    exceptions: List[Dict[str, Any]] = []
    sites = read_input_sites(input_path, aliases)

    if sites.empty:
        raise SystemExit(f"No site rows found in {input_path}")

    # Normalize tickets + numeric fields
    tpds, tpms, tpys = [], [], []
    for _, row in sites.iterrows():
        d, m, y = normalize_tickets(row, working_days, months)
        tpds.append(d)
        tpms.append(m)
        tpys.append(y)
        if d == 0 and to_float(row.get("users"), 0) == 0 and to_float(row.get("seat_count"), 0) == 0:
            exceptions.append(
                _exc(
                    len(exceptions) + 1,
                    row.get("site_id"),
                    row.get("site_name"),
                    "missing_demand",
                    "medium",
                    "No tickets/users/seats provided.",
                    "Add ticket volume or user/seat counts.",
                )
            )
    sites["tickets_per_day"] = tpds
    sites["tickets_per_month"] = tpms
    sites["tickets_per_year"] = tpys
    sites["users"] = sites["users"].apply(lambda v: to_float(v, 0.0))
    sites["seat_count"] = sites["seat_count"].apply(lambda v: to_float(v, 0.0))

    # Ensure site_id
    for i, row in sites.iterrows():
        if pd.isna(row.get("site_id")) or str(row.get("site_id")).strip() == "":
            sites.at[i, "site_id"] = f"SITE-{i+1:04d}"

    if skip_geocode:
        if "geocode_status" not in sites.columns:
            sites["geocode_status"] = sites.apply(
                lambda r: "provided"
                if pd.notna(r.get("latitude")) and pd.notna(r.get("longitude"))
                else "skipped",
                axis=1,
            )
        for i, row in sites.iterrows():
            if pd.isna(row.get("latitude")) or pd.isna(row.get("longitude")):
                exceptions.append(
                    _exc(
                        len(exceptions) + 1,
                        row.get("site_id"),
                        row.get("site_name"),
                        "geocode_skipped",
                        "low",
                        "Geocoding skipped; coordinates missing.",
                        "Rerun without --skip-geocode or supply lat/lon.",
                    )
                )
    else:
        sites = geocode_sites(sites, exceptions)

    mapped = assign_roles(sites, rules, exceptions)
    summary = build_summary(mapped, rules)
    cost_load = build_cost_model_load(mapped)

    finished = datetime.now(timezone.utc)
    role_counts = mapped["mapped_role"].value_counts().to_dict()
    metadata = [
        ("operator_name", operator_name),
        ("started_at_utc", started.isoformat(timespec="seconds")),
        ("finished_at_utc", finished.isoformat(timespec="seconds")),
        ("input_file", str(input_path)),
        ("output_file", str(output_path)),
        ("rules_file", str(rules_path)),
        ("aliases_file", str(aliases_path)),
        ("rules_version", rules.get("rules_version", "1.0.0")),
        ("site_count", len(mapped)),
        ("staffed_count", int(role_counts.get("Staffed", 0))),
        ("local_count", int(role_counts.get("Local", 0))),
        ("remote_count", int(role_counts.get("Remote", 0))),
        ("exception_count", len(exceptions)),
        ("hub_threshold_tickets_per_day", rules.get("hub_threshold_tickets_per_day")),
        ("remote_threshold_tickets_per_day", rules.get("remote_threshold_tickets_per_day")),
        ("cluster_radius_miles", rules.get("cluster_radius_miles")),
        ("max_drive_minutes", rules.get("max_drive_minutes")),
        ("mapper_version", "1.0.0"),
    ]

    write_output_workbook(
        output_path,
        input_path,
        mapped,
        summary,
        cost_load,
        exceptions,
        metadata,
        aliases,
        rules,
    )

    return {
        "site_count": len(mapped),
        "staffed": int(role_counts.get("Staffed", 0)),
        "local": int(role_counts.get("Local", 0)),
        "remote": int(role_counts.get("Remote", 0)),
        "exceptions": len(exceptions),
        "output": str(output_path),
    }


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Site Services VC mapper")
    p.add_argument("--input", required=True, help="Input workbook path")
    p.add_argument("--output", required=True, help="Output workbook path")
    p.add_argument("--operator-name", required=True, help="Operator name for Run_Metadata")
    p.add_argument("--rules", default=None, help="Path to default_rules.yaml")
    p.add_argument("--aliases", default=None, help="Path to column_aliases.yaml")
    p.add_argument(
        "--skip-geocode",
        action="store_true",
        help="Do not call Nominatim (use provided coordinates only)",
    )
    return p.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv)
    result = run_mapper(
        input_path=Path(args.input).expanduser().resolve(),
        output_path=Path(args.output).expanduser().resolve(),
        operator_name=args.operator_name,
        rules_path=Path(args.rules).expanduser().resolve() if args.rules else None,
        aliases_path=Path(args.aliases).expanduser().resolve() if args.aliases else None,
        skip_geocode=args.skip_geocode,
    )
    print(
        f"Mapped {result['site_count']} sites → "
        f"Staffed={result['staffed']} Local={result['local']} Remote={result['remote']} "
        f"Exceptions={result['exceptions']}"
    )
    print(f"Wrote {result['output']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
