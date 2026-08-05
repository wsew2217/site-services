from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


NAVY = "13294B"
STEEL = "54607A"
LT = "D9E1F2"
WHITE = "FFFFFF"
ACC = "149887"
thin = Side(style="thin", color="BFBFBF")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)


def _header(cell, fill=NAVY):
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BD


def _cell(cell, bold=False, fill=None):
    cell.font = Font(bold=bold, size=10)
    cell.border = BD
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def _write_df(ws, df: pd.DataFrame, start_row: int = 3):
    if df is None or len(df) == 0:
        ws.cell(start_row, 1, "(no rows)")
        return start_row
    for j, h in enumerate(df.columns, 1):
        _header(ws.cell(start_row, j, h))
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = ""
            _cell(ws.cell(i, j, v))
    for j in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = f"A{start_row + 1}"
    return start_row + 1 + len(df)


def write_deal_output(
    path: Path,
    sites: pd.DataFrame,
    vc_plan: pd.DataFrame,
    staffing_summary: Dict[str, Any],
    dispatch_detail: pd.DataFrame,
    dispatch_summary: Dict[str, Any],
    cost: Dict[str, Any],
    ops: Dict[str, Any],
    issues: List[Dict[str, Any]],
    settings: Dict[str, Any],
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Dashboard
    db = wb.active
    db.title = "Dashboard"
    db.sheet_view.showGridLines = False
    db["B1"] = "Site Services Deal Engine — Output"
    db["B1"].font = Font(bold=True, size=18, color=NAVY)
    db["B2"] = "Generated from intake. No customer names. Ratios are outputs/guardrails."
    db["B2"].font = Font(italic=True, size=10, color=STEEL)

    ov = staffing_summary["overlays"]
    kpis = [
        ("Sites", f"{staffing_summary['sites']:,}"),
        ("Tickets / year", f"{staffing_summary['tickets_yr']:,}"),
        ("Users (eff.)", f"{staffing_summary['users']:,}"),
        ("Countries", str(staffing_summary["countries"])),
        ("Virtual campuses", str(staffing_summary["vcs"])),
        ("Field technicians", str(ov["field_techs"])),
        ("Leads", str(ov["leads"])),
        ("Managers", str(ov["managers"])),
        ("Tech bars", str(ov["tech_bars"])),
        ("Depot techs", str(ov["depot_techs"])),
        ("Total people", str(ov["total_people"])),
        ("Day-one annual $", f"${cost['scenarios'][1]['c']['total']:,.0f}"),
    ]
    for i, (k, v) in enumerate(kpis):
        rr = 4 + (i // 4) * 3
        cc = 2 + (i % 4) * 3
        lab = db.cell(rr, cc, k)
        lab.font = Font(bold=True, size=9, color=WHITE)
        lab.fill = PatternFill("solid", fgColor=STEEL)
        db.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + 1)
        val = db.cell(rr + 1, cc, v)
        val.font = Font(bold=True, size=16, color=NAVY)
        val.fill = PatternFill("solid", fgColor=LT)
        db.merge_cells(start_row=rr + 1, start_column=cc, end_row=rr + 2, end_column=cc + 1)

    db.cell(14, 2, "Day-one staffing is throughput-derived from sites. Roadmap ratios are commercial guardrails only.")
    db.cell(15, 2, f"Implied day-one ratio: {cost['scenarios'][1]['c']['ratioImplied']}:1 users per field tech.")
    for col in range(1, 14):
        db.column_dimensions[get_column_letter(col)].width = 12

    # VC Plan
    p = wb.create_sheet("VC Plan")
    p["A1"] = "Virtual Campus Plan"
    p["A1"].font = Font(bold=True, size=14, color=NAVY)
    cols = [
        "Region", "VC", "Hub", "Countries", "StaffedSites", "LocalSites", "RemoteSites",
        "Served", "WorkDay", "TicketsYr", "Users", "OnSite", "Remote", "Team",
        "CampusMode", "Customs", "Staffed", "TechBar", "DepotTechs", "CriticalSites", "LaborCost",
    ]
    show = vc_plan[[c for c in cols if c in vc_plan.columns]].copy() if len(vc_plan) else pd.DataFrame(columns=cols)
    _write_df(p, show)

    # Site Roster
    sr = wb.create_sheet("Site Roster")
    sr["A1"] = "Site Roster"
    sr["A1"].font = Font(bold=True, size=14, color=NAVY)
    site_cols = [
        "Site", "Country", "Region", "City", "VCkey", "Role", "CoverageClassResolved",
        "TicketsYr", "TicketSource", "DemandConfidence", "UsersEff", "Devices",
        "Latitude", "Longitude", "DistKm", "CustomsFlag", "CustomerStaffedFlag",
        "Critical24x7Flag", "Notes",
    ]
    site_show = sites[[c for c in site_cols if c in sites.columns]].copy()
    _write_df(sr, site_show)

    # Staffing
    st = wb.create_sheet("Staffing")
    st["A1"] = "Staffing Summary"
    st["A1"].font = Font(bold=True, size=14, color=NAVY)
    st["A3"] = "Field technicians"
    st["B3"] = ov["field_techs"]
    st["A4"] = "Leads"
    st["B4"] = ov["leads"]
    st["A5"] = "Managers"
    st["B5"] = ov["managers"]
    st["A6"] = "Tech bars"
    st["B6"] = ov["tech_bars"]
    st["A7"] = "Depot techs"
    st["B7"] = ov["depot_techs"]
    st["A8"] = "Virtual techs"
    st["B8"] = ov["virtual_techs"]
    st["A9"] = "Total people"
    st["B9"] = ov["total_people"]
    st["A11"] = "Region rollup"
    st["A11"].font = Font(bold=True, color=NAVY)
    by = staffing_summary.get("by_region")
    if isinstance(by, pd.DataFrame):
        _write_df(st, by, start_row=12)

    # Overlays
    ol = wb.create_sheet("Overlays")
    ol["A1"] = "Overlays & coverage classes"
    ol["A1"].font = Font(bold=True, size=14, color=NAVY)
    ol["A3"] = "Coverage class counts"
    if "CoverageClassResolved" in sites.columns:
        counts = sites["CoverageClassResolved"].value_counts().reset_index()
        counts.columns = ["CoverageClass", "Sites"]
        _write_df(ol, counts, 4)
    ol["A12"] = "Role counts"
    if "Role" in sites.columns:
        roles = sites["Role"].value_counts().reset_index()
        roles.columns = ["Role", "Sites"]
        _write_df(ol, roles, 13)

    # Dispatch
    di = wb.create_sheet("Dispatch")
    di["A1"] = "Dispatch & shipment (variable)"
    di["A1"].font = Font(bold=True, size=14, color=NAVY)
    di["A2"] = (
        f"Events — dispatch {dispatch_summary.get('dispatch_events', 0):,.0f} / "
        f"shipments {dispatch_summary.get('shipment_events', 0):,.0f} | "
        f"Cost ${dispatch_summary.get('total_variable', 0):,.0f}"
    )
    _write_df(di, dispatch_detail)

    # Cost
    co = wb.create_sheet("Cost")
    co["A1"] = "Cost composition by scenario"
    co["A1"].font = Font(bold=True, size=14, color=NAVY)
    co["A2"] = "Modern rate placeholders from Settings. Legacy 2012-era models inform structure only — refresh rates per deal."
    headers = [
        "Scenario", "Techs", "Leads", "Mgrs", "Implied ratio", "Tech $", "Mgmt $",
        "Ship $", "Dispatch $", "OEM $", "Depot $", "Stock $", "Refresh $", "Total $",
    ]
    for j, h in enumerate(headers, 1):
        _header(co.cell(4, j, h))
    for i, item in enumerate(cost["scenarios"], 5):
        s, c = item["sc"], item["c"]
        vals = [
            s["name"], c["techs"], c["leads"], c["mgrs"], f"{c['ratioImplied']}:1",
            c["techCost"], c["mgmtCost"], c["shipCost"], c["dispCost"], c["oemCost"],
            c["depotCost"], c["stockCost"], c["rCost"], c["total"],
        ]
        for j, v in enumerate(vals, 1):
            _cell(co.cell(i, j, v), bold=(j == 1 or j == len(vals)))
            if j >= 6 and isinstance(v, (int, float)):
                co.cell(i, j).number_format = "$#,##0"
    for j in range(1, 15):
        co.column_dimensions[get_column_letter(j)].width = 12

    # Scenarios
    sc = wb.create_sheet("Scenarios")
    sc["A1"] = "Scenario ladder"
    sc["A1"].font = Font(bold=True, size=14, color=NAVY)
    for j, h in enumerate(["Key", "Name", "Ratio target", "Deflection %", "Zero-touch %", "Year", "Goal"], 1):
        _header(sc.cell(3, j, h))
    for i, item in enumerate(cost["scenarios"], 4):
        s = item["sc"]
        for j, v in enumerate([
            s["key"], s["name"], s.get("ratio") or "derived", s.get("defl"), s.get("zt"),
            s.get("year") if s.get("year") is not None else "today", "Y" if s.get("goal") else "",
        ], 1):
            _cell(sc.cell(i, j, v))

    # Cash
    ca = wb.create_sheet("Cash")
    ca["A1"] = "Transition cash (7-year)"
    ca["A1"].font = Font(bold=True, size=14, color=NAVY)
    cash = cost["cash"]
    ca["A3"] = "Total one-time"
    ca["B3"] = cash["total_one_time"]
    ca["B3"].number_format = "$#,##0"
    ca["A4"] = "Payback year"
    ca["B4"] = cash["payback_year"] if cash["payback_year"] is not None else ">7"
    ca["A5"] = "Year-7 net vs current"
    ca["B5"] = cash["year7_net"]
    ca["B5"].number_format = "$#,##0"
    for j, h in enumerate(["Step", "Year", "Exits", "Hires", "One-time", "New annual"], 1):
        _header(ca.cell(7, j, h))
    for i, step in enumerate(cost["steps"], 8):
        for j, v in enumerate([
            f"{step['from']} → {step['to']}", step["year"], step["exits"], step["hires"],
            step["oneOff"], step["annual"],
        ], 1):
            _cell(ca.cell(i, j, v))
            if j >= 5:
                ca.cell(i, j).number_format = "$#,##0"

    # Ops
    op = wb.create_sheet("Ops")
    op["A1"] = "Ops design QA (Reference Deal F patterns)"
    op["A1"].font = Font(bold=True, size=14, color=NAVY)
    summ = ops["summary"]
    r = 3
    for k, v in summ.items():
        if isinstance(v, dict):
            continue
        op.cell(r, 1, k)
        op.cell(r, 2, str(v))
        r += 1
    op.cell(r + 1, 1, "Staffing gap by VC")
    op.cell(r + 1, 1).font = Font(bold=True, color=NAVY)
    _write_df(op, ops["staffing_gap_by_vc"], r + 2)
    note_row = r + 4 + max(len(ops["staffing_gap_by_vc"]), 1)
    for i, note in enumerate(ops.get("notes", [])):
        op.cell(note_row + i, 1, f"• {note}")

    # Data Quality
    dq = wb.create_sheet("Data Quality")
    dq["A1"] = "Data quality flags"
    dq["A1"].font = Font(bold=True, size=14, color=NAVY)
    iss = pd.DataFrame(issues) if issues else pd.DataFrame(columns=["site", "level", "message"])
    _write_df(dq, iss)
    if "DemandConfidence" in sites.columns:
        conf = sites["DemandConfidence"].value_counts().reset_index()
        conf.columns = ["DemandConfidence", "Sites"]
        dq.cell(3 + len(iss) + 2, 1, "Demand confidence")
        _write_df(dq, conf, 3 + len(iss) + 3)

    # Assumptions
    ass = wb.create_sheet("Assumptions")
    ass["A1"] = "Settings used for this run"
    ass["A1"].font = Font(bold=True, size=14, color=NAVY)
    ass["A2"] = "Dollar rates are modern placeholders — refresh per deal. Older PFS workbooks inform structure only."
    for j, h in enumerate(["Parameter", "Value"], 1):
        _header(ass.cell(4, j, h))
    for i, (k, v) in enumerate(sorted(settings.items()), 5):
        _cell(ass.cell(i, 1, k))
        _cell(ass.cell(i, 2, v if not isinstance(v, (dict, list)) else json.dumps(v)))
    ass.column_dimensions["A"].width = 40
    ass.column_dimensions["B"].width = 20

    wb.save(path)
    return path


def write_web_summary(path: Path, staffing_summary: Dict[str, Any], cost: Dict[str, Any], settings: Dict[str, Any]) -> Path:
    """JSON for the live cost-model page to consume."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    web = cost.get("web_inputs", {})
    payload = {
        "source": "site-services-deal-engine",
        "version": "1.0.0",
        "note": "Campuses and day-one techs are engine-derived. Refresh labor/unit rates per deal.",
        "sites": web.get("sites"),
        "users": web.get("users"),
        "tickets_yr": web.get("tickets_yr"),
        "campuses": web.get("campuses"),
        "day1_techs": web.get("day1_techs"),
        "tpu": web.get("tpu"),
        "overlays": staffing_summary.get("overlays"),
        "scenarios": [
            {
                "key": s["sc"]["key"],
                "name": s["sc"]["name"],
                "ratio": s["sc"].get("ratio"),
                "defl": s["sc"].get("defl"),
                "zt": s["sc"].get("zt"),
                "year": s["sc"].get("year"),
                "techs": s["c"]["techs"],
                "total": s["c"]["total"],
                "ratioImplied": s["c"]["ratioImplied"],
            }
            for s in cost["scenarios"]
        ],
        "settings_subset": {
            "working_days": settings.get("working_days"),
            "utilization": settings.get("utilization"),
            "team_floor": settings.get("team_floor"),
            "physical_touch_share": settings.get("physical_touch_share"),
            "day_one_reach": settings.get("day_one_reach"),
        },
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return path
