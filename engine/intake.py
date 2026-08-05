from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import default_settings, merge_settings


SITE_HEADERS = [
    "Site", "Country", "Region", "City", "State", "Address", "PostalCode",
    "Latitude", "Longitude", "TicketsYr", "Users", "Seats", "Devices",
    "DepotHub", "Customs", "CustomerStaffed", "CoverageClass", "Critical24x7",
    "SpaceNeeded", "Zone", "Notes",
]

COUNTRY_HEADERS = [
    "Country", "Region", "Language", "CustomsLocked", "Embargoed",
    "DispatchPartner", "Notes",
]


def write_intake_template(path: Path) -> Path:
    """Create a blank editable intake workbook."""
    path = Path(path)
    wb = Workbook()

    # Sites
    ws = wb.active
    ws.title = "Sites"
    for j, h in enumerate(SITE_HEADERS, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="13294B")
    sample = [
        ["Site Alpha HQ", "USA", "NAM", "Austin", "TX", "", "", 30.27, -97.74, 4200, 1800, "", "", "USA-Central", "", "", "", "", "Y", "Campus", ""],
        ["Site Alpha East", "USA", "NAM", "Dallas", "TX", "", "", 32.78, -96.80, 900, 400, "", "", "USA-Central", "", "", "", "", "", "Metro", ""],
        ["Site Beta Hub", "Germany", "EMEA", "Frankfurt", "", "", "", 50.11, 8.68, 3100, 1200, "", "", "", "Y", "", "", "", "", "Campus", ""],
        ["Site Beta Spoke", "Germany", "EMEA", "Mainz", "", "", "", 49.99, 8.27, 180, 80, "", "", "", "", "", "", "", "", "", ""],
        ["Site Gamma Remote", "Brazil", "LATAM", "Manaus", "", "", "", "", "", "", 60, "", "", "", "Y", "", "Dispatch", "", "", "Remote", "No coords — remote/uncertain"],
        ["Site Delta Devices Only", "Singapore", "APJC", "Singapore", "", "", "", 1.29, 103.85, "", "", "", 500, "", "", "", "", "", "", "", "Demand from devices"],
    ]
    for i, row in enumerate(sample, 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
    from openpyxl.utils import get_column_letter
    widths = [22, 12, 8, 14, 8, 18, 10, 10, 10, 10, 8, 8, 8, 12, 9, 12, 14, 10, 10, 10, 28]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # Settings
    st = wb.create_sheet("Settings")
    st["A1"] = "Parameter"
    st["B1"] = "Value"
    st["C1"] = "Notes"
    for c in ("A1", "B1", "C1"):
        st[c].font = Font(bold=True, color="FFFFFF")
        st[c].fill = PatternFill("solid", fgColor="13294B")
    notes = {
        "working_days": "Standard working days / year",
        "drive_radius_km": "Local/Staffed catchment radius",
        "incident_share": "0-1; calibrate from ticket extract when available",
        "include_cost": "1 = include labor/cost tabs",
        "estate_tickets_per_user": "Used when tickets missing but users/seats present",
        "estate_tickets_per_device": "Used when only device counts present",
        "placeholder_tickets_per_site": "Last-resort demand when no volume fields",
        "team_floor": "Minimum techs per VC for resilience",
        "single_tech_threshold": "Tickets/day below which 1-tech exception allowed",
    }
    for i, (k, v) in enumerate(default_settings().items(), 2):
        st.cell(i, 1, k)
        st.cell(i, 2, v)
        st.cell(i, 3, notes.get(k, ""))
    st.column_dimensions["A"].width = 36
    st.column_dimensions["B"].width = 14
    st.column_dimensions["C"].width = 48

    # CountryPolicy
    cp = wb.create_sheet("CountryPolicy")
    for j, h in enumerate(COUNTRY_HEADERS, 1):
        c = cp.cell(1, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="13294B")
    for i, row in enumerate([
        ["USA", "NAM", "English", "", "", "", ""],
        ["Germany", "EMEA", "German", "", "", "", ""],
        ["Brazil", "LATAM", "Portuguese", "Y", "", "", "Domestic stock/staff when customs-locked"],
        ["Singapore", "APJC", "English", "", "", "", ""],
        ["United Kingdom", "EMEA", "English", "Y", "", "", ""],
        ["China", "APJC", "Mandarin", "Y", "", "", ""],
        ["Japan", "APJC", "Japanese", "Y", "", "", ""],
        ["Canada", "NAM", "English/French", "Y", "", "", ""],
        ["Australia", "APJC", "English", "Y", "", "", ""],
    ], 2):
        for j, v in enumerate(row, 1):
            cp.cell(i, j, v)
    for j, w in enumerate([18, 8, 16, 12, 10, 16, 40], 1):
        cp.column_dimensions[get_column_letter(j)].width = w

    readme = wb.create_sheet("ReadMe", 0)
    readme["A1"] = "Site Services Deal Engine — Intake Template"
    readme["A1"].font = Font(bold=True, size=16, color="13294B")
    lines = [
        "",
        "1. Fill the Sites tab. TicketsYr and/or Users/Seats and/or Devices are all accepted.",
        "2. Leave DepotHub blank to default one virtual campus per Country.",
        "3. Tweak Settings (yellow conceptually — all Value cells are inputs).",
        "4. Optional CountryPolicy for customs/language/partner notes.",
        "5. Run:  python -m engine run path/to/this.xlsx -o Deal_Output.xlsx",
        "",
        "Demand resolution order: tickets → users/seats × rate → devices × rate → placeholder (flagged).",
        "Sites without coordinates are never auto-classified Local/Staffed by distance.",
        "No customer names belong in this template — use generic site labels.",
    ]
    for i, line in enumerate(lines, 2):
        readme.cell(i, 1, line)
        readme.cell(i, 1).alignment = Alignment(wrap_text=True)
    readme.column_dimensions["A"].width = 100

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_intake(path: Path) -> Tuple[pd.DataFrame, Dict[str, Any], pd.DataFrame]:
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    settings_raw: Dict[str, Any] = {}
    if "Settings" in wb.sheetnames:
        for row in wb["Settings"].iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            settings_raw[str(row[0]).strip()] = row[1]
    settings = merge_settings(settings_raw)

    sites = pd.read_excel(path, sheet_name="Sites")
    if "CountryPolicy" in wb.sheetnames:
        policy = pd.read_excel(path, sheet_name="CountryPolicy")
    else:
        policy = pd.DataFrame(columns=COUNTRY_HEADERS)
    wb.close()
    return sites, settings, policy
