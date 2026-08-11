#!/usr/bin/env python3
"""Normalize a customer address-request workbook into VC_Site_Input_Template format.

LOCAL USE ONLY — do not commit customer-named outputs to git.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from vc_mapper import INPUT_HEADERS, style_header

HEADER_FILL = PatternFill("solid", fgColor="13294B")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _cell_str(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def read_address_request(path: Path) -> pd.DataFrame:
    """Parse the Address Request sheet (header row is not row 1)."""
    raw = pd.read_excel(path, sheet_name="Address Request", header=None)
    # Find header row containing 'Country' and 'Site name'
    header_idx = None
    for i, row in raw.iterrows():
        vals = [_cell_str(v).lower() for v in row.tolist()]
        if any("country" == v for v in vals) and any("site name" in v for v in vals):
            header_idx = i
            break
    if header_idx is None:
        raise SystemExit("Could not find header row on Address Request sheet")

    headers = [_cell_str(v) for v in raw.iloc[header_idx].tolist()]
    data = raw.iloc[header_idx + 1 :].copy()
    data.columns = headers
    # Drop empty rows
    data = data.dropna(how="all")
    return data


def normalize(df: pd.DataFrame) -> pd.DataFrame:
    colmap = {}
    for c in df.columns:
        cl = _cell_str(c).lower()
        if cl in ("#", "no", "num"):
            colmap[c] = "_num"
        elif "country" == cl:
            colmap[c] = "country"
        elif "site name" in cl:
            colmap[c] = "site_name"
        elif cl.startswith("city on"):
            colmap[c] = "city_on_file"
        elif "tickets" in cl:
            colmap[c] = "tickets_per_year"
        elif "street" in cl or cl == "address":
            colmap[c] = "address_full"
        elif "zip" in cl or "postal" in cl:
            colmap[c] = "postal_code"
        elif "city (confirm)" in cl or cl == "city":
            colmap[c] = "city_confirm"
        elif "users" in cl:
            colmap[c] = "users"
        elif "notes" in cl:
            colmap[c] = "notes"
    df = df.rename(columns=colmap)

    rows = []
    for i, r in df.iterrows():
        site_name = _cell_str(r.get("site_name"))
        if not site_name:
            continue
        city = _cell_str(r.get("city_confirm")) or _cell_str(r.get("city_on_file"))
        # Some rows incorrectly put country in city confirm
        country = _cell_str(r.get("country"))
        if city.lower() == country.lower():
            city = _cell_str(r.get("city_on_file"))
        # Clean city values that embed multiple cities
        if "/" in city:
            city = city.split("/")[0].strip()
        tickets = r.get("tickets_per_year")
        try:
            tickets = float(tickets) if tickets is not None and str(tickets).strip() != "" else 0.0
        except (TypeError, ValueError):
            tickets = 0.0
        users = r.get("users")
        try:
            users = float(users) if users is not None and str(users).strip() != "" else 0.0
        except (TypeError, ValueError):
            users = 0.0
        postal = _cell_str(r.get("postal_code"))
        if postal.lower() in {"twin", "n/a", "na", "#n/a", "none"}:
            postal = ""
        # Algeria extract used country dial code as ZIP
        if country.lower() == "algeria" and postal == "213":
            postal = ""
        address = _cell_str(r.get("address_full"))
        if address.lower().startswith("n/a"):
            address = ""
        site_id = f"PILOT-{i:04d}" if not _cell_str(r.get("_num")) else f"PILOT-{_cell_str(r.get('_num')).zfill(4)}"
        rows.append(
            {
                "site_id": site_id,
                "site_name": site_name,
                "address_full": address,
                "city": city,
                "state_province": "",
                "postal_code": postal,
                "country": country,
                "users": users,
                "seat_count": "",
                "tickets_per_day": "",
                "tickets_per_month": "",
                "tickets_per_year": tickets,
                "region": "",
                "district": "",
                "customer_site_type": "",
                "notes": _cell_str(r.get("notes")),
            }
        )
    return pd.DataFrame(rows, columns=INPUT_HEADERS)


def write_input_workbook(df: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "VC_Site_Input_Template"
    for j, h in enumerate(INPUT_HEADERS, 1):
        cell = ws.cell(1, j, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, h in enumerate(INPUT_HEADERS, 1):
            ws.cell(i, j, getattr(row, h))
    end_row = max(2, len(df) + 1)
    end_col = get_column_letter(len(INPUT_HEADERS))
    table = Table(displayName="tbl_vc_site_input", ref=f"A1:{end_col}{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False,
        showLastColumn=False, showColumnStripes=False,
    )
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{end_col}{end_row}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--input", required=True)
    p.add_argument("--output", required=True)
    args = p.parse_args()
    raw = read_address_request(Path(args.input).expanduser())
    norm = normalize(raw)
    out = Path(args.output).expanduser()
    write_input_workbook(norm, out)
    print(f"Normalized {len(norm)} sites → {out}")


if __name__ == "__main__":
    main()
