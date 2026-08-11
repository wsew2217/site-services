#!/usr/bin/env python3
"""Bridge Cost_Model_Load → engine Sites intake + rough cost-model summary JSON.

Does not invent commercial quotes. day1_techs come from mapper hints / Staffed roles.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Dict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ENGINE_SITE_HEADERS = [
    "Site",
    "Country",
    "Region",
    "City",
    "State",
    "Address",
    "PostalCode",
    "Latitude",
    "Longitude",
    "TicketsYr",
    "Users",
    "Seats",
    "Devices",
    "DepotHub",
    "Customs",
    "CustomerStaffed",
    "CoverageClass",
    "Critical24x7",
    "SpaceNeeded",
    "Zone",
    "Notes",
]


def load_cost_model_sheet(path: Path) -> pd.DataFrame:
    xl = pd.ExcelFile(path)
    sheet = "Cost_Model_Load" if "Cost_Model_Load" in xl.sheet_names else xl.sheet_names[0]
    return pd.read_excel(path, sheet_name=sheet)


def to_engine_sites(df: pd.DataFrame) -> pd.DataFrame:
    out = pd.DataFrame(columns=ENGINE_SITE_HEADERS)
    for h in ENGINE_SITE_HEADERS:
        out[h] = df[h] if h in df.columns else ""
    return out


def summarize_for_cost_page(df: pd.DataFrame) -> Dict[str, Any]:
    roles = df["mapped_role"].astype(str) if "mapped_role" in df.columns else pd.Series([""] * len(df))
    staffed = int((roles == "Staffed").sum())
    local = int((roles == "Local").sum())
    remote = int((roles == "Remote").sum())
    users = float(pd.to_numeric(df.get("Users"), errors="coerce").fillna(0).sum())
    tickets_yr = float(pd.to_numeric(df.get("TicketsYr"), errors="coerce").fillna(0).sum())
    if "day1_techs_hint" in df.columns:
        day1 = float(pd.to_numeric(df["day1_techs_hint"], errors="coerce").fillna(0).sum())
    else:
        day1 = float(staffed * 2)  # floor-ish fallback; flag as rough
    campuses = max(staffed, 1) if len(df) else 0
    tpu = (tickets_yr / users) if users else 0.0
    return {
        "source": "vc-mapper-cost-model-bridge",
        "version": "1.0.0",
        "note": (
            "Rough bridge from Cost_Model_Load — not a full engine run. "
            "Prefer python -m engine run on the Sites intake for production summaries."
        ),
        "sites": int(len(df)),
        "users": round(users, 0),
        "tickets_yr": round(tickets_yr, 0),
        "campuses": int(campuses),
        "day1_techs": round(day1, 1),
        "tpu": round(tpu, 3),
        "role_split": {"Staffed": staffed, "Local": local, "Remote": remote},
        "overlays": {
            "field_techs": round(day1, 1),
            "leads": 0,
            "managers": 0,
            "tech_bars": 0,
            "depot_techs": 0,
            "virtual_techs": 0,
            "total_people": round(day1, 1),
        },
    }


def write_sites_xlsx(df: pd.DataFrame, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sites"
    fill = PatternFill("solid", fgColor="13294B")
    font = Font(bold=True, color="FFFFFF")
    for j, h in enumerate(ENGINE_SITE_HEADERS, 1):
        c = ws.cell(1, j, h)
        c.fill = fill
        c.font = font
    for i, row in enumerate(df.itertuples(index=False), 2):
        mapping = row._asdict() if hasattr(row, "_asdict") else None
        for j, h in enumerate(ENGINE_SITE_HEADERS, 1):
            val = mapping.get(h) if mapping else df.iloc[i - 2].get(h)
            if pd.isna(val):
                val = None
            ws.cell(i, j, val)
    for j in range(1, len(ENGINE_SITE_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> None:
    p = argparse.ArgumentParser(description="Cost_Model_Load → engine Sites / summary JSON")
    p.add_argument("--input", required=True, help="Mapped workbook with Cost_Model_Load")
    p.add_argument("--sites-out", default="engine_sites_from_vc.xlsx")
    p.add_argument("--json-out", default="vc_bridge_summary.json")
    args = p.parse_args()

    src = Path(args.input)
    df = load_cost_model_sheet(src)
    sites = to_engine_sites(df)
    sites_path = write_sites_xlsx(sites, Path(args.sites_out))
    summary = summarize_for_cost_page(df)
    jpath = Path(args.json_out)
    jpath.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(f"Wrote Sites intake: {sites_path}")
    print(f"Wrote rough summary JSON: {jpath}")
    print(
        f"sites={summary['sites']} campuses={summary['campuses']} "
        f"day1_techs≈{summary['day1_techs']} roles={summary['role_split']}"
    )


if __name__ == "__main__":
    main()
