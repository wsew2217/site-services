#!/usr/bin/env python3
"""Create vc_template_pack_v1_0.xlsx with openpyxl (Wave 1)."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from vc_mapper import (
    ALIAS_MAP_HEADERS,
    COST_MODEL_HEADERS,
    EXCEPTIONS_HEADERS,
    INPUT_HEADERS,
    MAPPED_HEADERS,
    RUN_METADATA_HEADERS,
    SHEET_ORDER,
    SUMMARY_HEADERS,
    build_alias_map_rows,
    load_yaml,
    style_header,
    write_instructions,
)

HEADER_FILL = PatternFill("solid", fgColor="13294B")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def apply_sheet_standards(ws, n_cols: int, n_rows: int = 1) -> None:
    ws.freeze_panes = "A2"
    end = get_column_letter(n_cols)
    ws.auto_filter.ref = f"A1:{end}{max(n_rows, 1)}"
    for col in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def write_headers(ws, headers) -> None:
    for j, h in enumerate(headers, 1):
        cell = ws.cell(1, j, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


# Generic starter rows — Reference Deal A style (no real customer names)
STARTER_ROWS = [
    {
        "site_id": "RDA-ATL-01",
        "site_name": "Reference Deal A — Atlanta HQ",
        "address_full": "1 Peachtree Center",
        "city": "Atlanta",
        "state_province": "GA",
        "postal_code": "30303",
        "country": "USA",
        "users": 420,
        "seat_count": 400,
        "tickets_per_day": 4.2,
        "tickets_per_month": "",
        "tickets_per_year": "",
        "region": "NAM",
        "district": "Southeast",
        "customer_site_type": "HQ",
        "notes": "Primary hub candidate",
        "latitude": 33.7590,
        "longitude": -84.3880,
    },
    {
        "site_id": "RDA-ATL-02",
        "site_name": "Reference Deal A — Midtown Satellite",
        "address_full": "1000 West Peachtree St",
        "city": "Atlanta",
        "state_province": "GA",
        "postal_code": "30309",
        "country": "USA",
        "users": 85,
        "seat_count": 80,
        "tickets_per_day": 0.9,
        "tickets_per_month": "",
        "tickets_per_year": "",
        "region": "NAM",
        "district": "Southeast",
        "customer_site_type": "Office",
        "notes": "Near Atlanta HQ — expect Local",
        "latitude": 33.7840,
        "longitude": -84.3900,
    },
    {
        "site_id": "RDA-CHI-01",
        "site_name": "Reference Deal A — Chicago Loop",
        "address_full": "233 S Wacker Dr",
        "city": "Chicago",
        "state_province": "IL",
        "postal_code": "60606",
        "country": "USA",
        "users": 210,
        "seat_count": 200,
        "tickets_per_day": 2.5,
        "tickets_per_month": "",
        "tickets_per_year": "",
        "region": "NAM",
        "district": "Central",
        "customer_site_type": "Regional",
        "notes": "Far from Atlanta; >=1.2 tpd → Staffed campus",
        "latitude": 41.8789,
        "longitude": -87.6359,
    },
    {
        "site_id": "RDA-BOI-01",
        "site_name": "Reference Deal A — Boise Remote",
        "address_full": "250 S 5th St",
        "city": "Boise",
        "state_province": "ID",
        "postal_code": "83702",
        "country": "USA",
        "users": 18,
        "seat_count": 16,
        "tickets_per_day": 0.3,
        "tickets_per_month": "",
        "tickets_per_year": "",
        "region": "NAM",
        "district": "Mountain",
        "customer_site_type": "Small office",
        "notes": "Far + <1.2 tpd → Remote / dispatch",
        "latitude": 43.6150,
        "longitude": -116.2023,
    },
    {
        "site_id": "RDA-LON-01",
        "site_name": "Reference Deal A — London Docklands",
        "address_full": "1 Canada Square",
        "city": "London",
        "state_province": "",
        "postal_code": "E14 5AB",
        "country": "United Kingdom",
        "users": 150,
        "seat_count": 140,
        "tickets_per_month": 55,
        "tickets_per_day": "",
        "tickets_per_year": "",
        "region": "EMEA",
        "district": "UK",
        "customer_site_type": "Regional",
        "notes": "Month-based tickets; EMEA campus seed",
        "latitude": 51.5048,
        "longitude": -0.0195,
    },
]


def build_template(path: Path, rules: dict, aliases: dict) -> Path:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    # Instructions
    ws_inst = wb.create_sheet("Instructions")
    write_instructions(ws_inst, rules)

    # Input — required headers exact; optional latitude/longitude appended for starter geocode skip
    input_headers_ext = INPUT_HEADERS + ["latitude", "longitude"]
    ws_in = wb.create_sheet("VC_Site_Input_Template")
    write_headers(ws_in, input_headers_ext)
    for i, row in enumerate(STARTER_ROWS, 2):
        for j, h in enumerate(input_headers_ext, 1):
            ws_in.cell(i, j, row.get(h, ""))
    end_row = 1 + len(STARTER_ROWS)
    end_col = get_column_letter(len(input_headers_ext))
    table = Table(displayName="tbl_vc_site_input", ref=f"A1:{end_col}{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws_in.add_table(table)
    apply_sheet_standards(ws_in, len(input_headers_ext), end_row)
    for col in range(1, len(input_headers_ext) + 1):
        ws_in.column_dimensions[get_column_letter(col)].width = 18

    # Output header tabs
    for name, headers in [
        ("VC_Mapped_Sites", MAPPED_HEADERS),
        ("VC_Summary", SUMMARY_HEADERS),
        ("Cost_Model_Load", COST_MODEL_HEADERS),
    ]:
        ws = wb.create_sheet(name)
        write_headers(ws, headers)
        apply_sheet_standards(ws, len(headers), 1)

    # Exceptions starter
    ws_exc = wb.create_sheet("Exceptions")
    write_headers(ws_exc, EXCEPTIONS_HEADERS)
    starter_exc = [
        "EXC-000",
        "",
        "",
        "template_placeholder",
        "info",
        "No exceptions yet — mapper populates this tab on each run.",
        "Review after first map.",
        "",
    ]
    for j, v in enumerate(starter_exc, 1):
        ws_exc.cell(2, j, v)
    apply_sheet_standards(ws_exc, len(EXCEPTIONS_HEADERS), 2)

    # Run_Metadata starter
    ws_meta = wb.create_sheet("Run_Metadata")
    write_headers(ws_meta, RUN_METADATA_HEADERS)
    meta_rows = [
        ("template_version", "1.0.0"),
        ("rules_version", rules.get("rules_version", "1.0.0")),
        ("operator_name", ""),
        ("last_run_utc", ""),
        ("notes", "Filled by vc_mapper.py on each run"),
    ]
    for i, (k, v) in enumerate(meta_rows, 2):
        ws_meta.cell(i, 1, k)
        ws_meta.cell(i, 2, v)
    apply_sheet_standards(ws_meta, 2, len(meta_rows) + 1)

    # Alias_Map starter from YAML
    ws_alias = wb.create_sheet("Alias_Map")
    write_headers(ws_alias, ALIAS_MAP_HEADERS)
    alias_df = build_alias_map_rows(aliases)
    for i, row in enumerate(alias_df.itertuples(index=False), 2):
        ws_alias.cell(i, 1, row.canonical_field)
        ws_alias.cell(i, 2, row.alias)
        ws_alias.cell(i, 3, row.notes)
    apply_sheet_standards(ws_alias, 3, max(2, len(alias_df) + 1))

    # Exact tab order
    wb._sheets = [wb[name] for name in SHEET_ORDER]  # noqa: SLF001

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> None:
    base = Path(__file__).resolve().parent
    rules = load_yaml(base / "default_rules.yaml")
    aliases = load_yaml(base / "column_aliases.yaml")
    out = base / "vc_template_pack_v1_0.xlsx"
    build_template(out, rules, aliases)

    # Publish copies
    dests = [
        Path.home() / "Desktop" / "Site Services Solutioning" / "01_Templates" / "vc_template_pack" / "vc_template_pack_v1_0.xlsx",
        Path("/Users/darrenreinhardt/GitHub/bv-vercel-site/samples/vc_template_pack_v1_0.xlsx"),
        Path("/Users/darrenreinhardt/GitHub/bv-vercel-site/tools/vc_mapper/vc_template_pack_v1_0.xlsx"),
    ]
    for d in dests:
        d.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(out, d)
        print(f"Wrote {d}")
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
